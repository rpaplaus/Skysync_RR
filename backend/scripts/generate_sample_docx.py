#!/usr/bin/env python3
"""
Gera um manuscrito DOCX de amostra para Skyhaven, contendo marcações de teste
para validar o parser docx_parser.py.
"""

from pathlib import Path
from docx import Document

DOCX_DIR = Path(__file__).resolve().parent.parent / "data" / "raw"
DOCX_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT_FILE = DOCX_DIR / "skyhaven.docx"

doc = Document()

# Título do Livro
doc.add_heading("Skyhaven: Crônicas do Éter", level=0)

# Capítulo 1
doc.add_heading("Capítulo 1: As Torres de Aço e Vapor", level=1)
doc.add_paragraph(
    "O vento sibilava entre as imensas vigas de sustentação de Skyhaven, fazendo ranger as placas de cobre corroídas pelo vapor salino das nuvens inferiores. A trezentos metros acima da terra devastada, a cidade flutuante resistia ao tempo e à gravidade com a teimosia dos antigos engenheiros."
)
doc.add_paragraph(
    "Kael ajustou seus óculos de proteção contra a fuligem. Olhando para o abismo sob a passarela de manutenção, ele sentia o pulsar rítmico das turbinas primárias — o coração adormecido da metrópole suspensa."
)
doc.add_paragraph(
    "Ao longe, um gramofone a válvula em uma das oficinas abertas começou a ecoar as notas metálicas e melancólicas de \"Aetherium Dreams\". O som parecia acariciar o ar rarefeito, trazendo lembranças de um mundo onde o solo ainda era habitável."
)
doc.add_paragraph(
    "De repente, um alarme de pressão estalou no console do painel três. As caldeiras rugiram com a fúria estrondosa de \"Engenhos a Vapor\", enquanto jatos brancos de vapor superaquecido forçavam os mecânicos a correrem pelos passadiços enferrujados. O tilintar frenético de ferramentas ecoava com a urgência da sobrevivência."
)
doc.add_paragraph(
    "— Bloqueiem a linha secundária antes que o condensador colapse! — gritou Kael, saltando sobre uma tubulação fumegante com a chave inglesa em punho."
)

# Capítulo 2
doc.add_heading("Capítulo 2: O Salão dos Navegadores", level=1)
doc.add_paragraph(
    "A poeira de carvão e a graxa dos níveis inferiores contrastavam brutalmente com a opulência do Distrito Superior. Lustres de cristal alimentados por filamentos de éter cintilavam sobre o piso de mogno polido do Grande Salão."
)
doc.add_paragraph(
    "Diplomatas e capitães de dirigíveis trajando veludo escuro e insígnias douradas brindavam com taças de champanhe enquanto uma orquestra de cordas executava a suave \"Valsa nas Alturas\". Os passos deslizavam pelo salão em perfeita harmonia, fingindo ignorar as tempestades elétricas no horizonte."
)
doc.add_paragraph(
    "Enquanto isso, nos conveses de vigia subterrâneos, a playlist do transmissor pirata continuou e \"Immigrant Song\" do Led Zeppelin explodiu baixo nas frequências de rádio, energizando a resistência dos operários da noite."
)

doc.save(str(OUTPUT_FILE))
print(f"Manuscrito de teste gerado com sucesso em: {OUTPUT_FILE}")
